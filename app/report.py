# -*- coding: utf-8 -*-
"""
bh_report - "КАМЕРАЛ ТЕКШИРУВЛАР" Excel hisoboti. HISOB-KITOB EXCEL FORMULALARIDA
(1.6.0, egasi talabi): ilova faqat ma'lumotni yozadi, kirim, chiqim, qoldiq,
sotish narxi va jamilarni Excelning o'zi hisoblaydi. Batafsil - build_formula_workbook.

Varaqlar:
    Ma'lumot    - QQS va ustama (sariq katak), yillar jamisi, oylar jadvali
    Фактуралар  - har faktura bir satr
    Кирим       - har faktura satri
    касса YYYY  - kassa cheklari (20 ustun) + tovar kodi
    YYYY ТХ     - tovar harakati (22 ustun, ikki qavatli sarlavha), hammasi formula
    Xatolar     - tekshiruv natijalari

Chiqish formati .xlsx.
"""

import os
import re
import datetime
from decimal import Decimal

import bh_config as C
import bh_db as DB
import bh_fifo as F


NUM_MONEY = '#,##0.00'
NUM_QTY = '#,##0.####'
NUM_DATE = 'DD.MM.YYYY'


class StyleBook:
    """
    NamedStyle keshi.

    openpyxl'da har `cell.font = ...` topshirig'i stilni xeshlab ish
    kitobining stil jadvalidan qidiradi. 10 000 satrli kassa varag'ida bu
    150 000 dan ortiq qidiruv - hisobotning eng sekin joyi aynan shu edi.

    NamedStyle bir marta ro'yxatdan o'tkaziladi, keyin `cell.style = nom`
    faqat tayyor indeksni ko'chiradi: ko'rinish bir xil, tezlik ~5 barobar.

    name(asos, raqam_formati, rang) -> stil nomi. Rang ("warn" / "err")
    asosning fonini almashtiradi, shuning uchun satrni bo'yash uchun
    ikkinchi marta aylanib chiqish shart emas.
    """

    def __init__(self, wb):
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

        self.wb = wb
        self._names = {}

        thin = Side(style="thin", color="9AA0A6")
        med = Side(style="medium", color="44546A")
        box = Border(left=thin, right=thin, top=thin, bottom=thin)

        self.font_cell = Font(name="Calibri", size=9)
        self.font_total = Font(name="Calibri", size=10, bold=True)

        self.tints = {
            "warn": PatternFill("solid", fgColor="FFF2CC"),
            "err": PatternFill("solid", fgColor="FCE4E4"),
        }

        self.bases = {
            "title": {"font": Font(name="Calibri", size=14, bold=True, color="1F3864"),
                      "align": Alignment(horizontal="center", vertical="center")},
            "hdr": {"font": Font(name="Calibri", size=9, bold=True, color="FFFFFF"),
                    "fill": PatternFill("solid", fgColor="44546A"),
                    "align": Alignment(horizontal="center", vertical="center",
                                       wrap_text=True),
                    "border": Border(left=thin, right=thin, top=med, bottom=thin)},
            "sub": {"font": Font(name="Calibri", size=9, bold=True, color="FFFFFF"),
                    "fill": PatternFill("solid", fgColor="5B7397"),
                    "align": Alignment(horizontal="center", vertical="center",
                                       wrap_text=True),
                    "border": Border(left=thin, right=thin, top=thin, bottom=med)},
            "cell": {"font": self.font_cell,
                     "align": Alignment(vertical="top", wrap_text=True),
                     "border": box},
            "num": {"font": self.font_cell,
                    "align": Alignment(horizontal="right", vertical="top"),
                    "border": box},
            "total": {"font": self.font_total,
                      "fill": PatternFill("solid", fgColor="D9E2F3"),
                      "align": Alignment(horizontal="right", vertical="center"),
                      "border": Border(left=thin, right=thin, top=med, bottom=med)},
        }

    def name(self, base, numfmt=None, tint=None):
        key = (base, numfmt or "", tint or "")
        nm = self._names.get(key)
        if nm is not None:
            return nm

        from openpyxl.styles import NamedStyle

        spec = self.bases[base]
        nm = "bh%d" % len(self._names)
        ns = NamedStyle(name=nm)
        if "font" in spec:
            ns.font = spec["font"]
        fill = self.tints[tint] if tint else spec.get("fill")
        if fill is not None:
            ns.fill = fill
        if "align" in spec:
            ns.alignment = spec["align"]
        if "border" in spec:
            ns.border = spec["border"]
        if numfmt:
            ns.number_format = numfmt
        self.wb.add_named_style(ns)
        self._names[key] = nm
        return nm



# ===========================================================================
# Xatolar varag'i (YANGI)
# ===========================================================================
def write_issues_sheet(wb, cx, years, S, owner_tin=None):
    from openpyxl.utils import get_column_letter

    ws = wb.create_sheet("Xatolar")
    heads = ["Yil", "Daraja", "Turi", "Tavsif", "Manba", "ID"]
    widths = [8, 16, 26, 96, 12, 8]
    for i, (h, w) in enumerate(zip(heads, widths)):
        ws.column_dimensions[get_column_letter(i + 1)].width = w
        ws.cell(row=1, column=i + 1, value=h).style = S.name("hdr")
    ws.row_dimensions[1].height = 22

    r = 2
    # Tashkilot tanlangan bo'lsa - faqat uning hujjatlariga oid yozuvlar
    q = ("SELECT i.* FROM issue i "
         "LEFT JOIN doc_line il ON i.ref_table = 'doc_line' AND il.id = i.ref_id "
         "LEFT JOIN document idoc ON idoc.id = CASE WHEN i.ref_table = 'document' "
         "     THEN i.ref_id ELSE il.document_id END "
         "WHERE i.resolved=0 AND (? IS NULL OR idoc.owner_tin = ? OR "
         "     (idoc.id IS NULL AND (i.code <> 'unmatched_sale' OR i.detail = ?)))")
    args = [owner_tin, owner_tin, owner_tin]
    if years:
        q += " AND (i.year IS NULL OR i.year IN (%s))" % ",".join("?" * len(years))
        args += list(years)
    q += (" ORDER BY CASE i.severity WHEN 'xato' THEN 0 WHEN 'ogohlantirish' "
          "THEN 1 ELSE 2 END, i.year, i.code")
    n = 0
    iss_plain = S.name("cell")
    iss_warn = S.name("cell", None, "warn")
    iss_err = S.name("cell", None, "err")
    for d in cx.execute(q, args):
        vals = [d["year"], d["severity"],
                C.ISSUE_TITLES.get(d["code"], d["code"]), d["message"],
                d["ref_table"], d["ref_id"]]
        if d["severity"] == C.SEVERITY_ERROR:
            sty = iss_err
        elif d["severity"] == C.SEVERITY_WARN:
            sty = iss_warn
        else:
            sty = iss_plain
        for i, v in enumerate(vals):
            ws.cell(row=r, column=i + 1, value=v).style = sty
        r += 1
        n += 1

    ws.freeze_panes = "A2"
    if n:
        ws.auto_filter.ref = "A1:F%d" % (r - 1)
    return n


# ===========================================================================
# Asosiy
# ===========================================================================
def generate(cx, out_path, years=None, owner_name=None, progress=None, owner_tin=None):
    """
    Hisobotni yozadi. (yo'l, statistika) qaytaradi.

    owner_tin - qaysi tashkilot (STIR) hisoboti. Berilmasa - hammasi (eski xulq).
    """
    years = years or DB.available_years(cx, owner_tin)
    years = sorted(y for y in years if y)
    if not years:
        raise ValueError("Hisobot uchun ma'lumot yo'q - avval fayllarni import qiling.")

    owner_name = owner_name or (DB.org_name(cx, owner_tin) if owner_tin else "") \
        or DB.get_setting(cx, "owner_name", "Ташкилот")

    # Hisobotdan oldin ombor majburan qayta hisoblanadi: qolda
    # boglangan sotuvlar ham qoldiqdan ayrilsin (Qayta hisoblash
    # tugmasi bosilmagan bolsa ham).
    import bh_matching as _M
    if progress:
        progress(0, 0, 'ombor yangilanmoqda')
    try:
        _eng = _M.MatchEngine(cx)
        # 1.3.0 gacha to'liq o'qilmagan chek fayllari (bir marta)
        F.reparse_old_checks(cx, _eng)
        _eng.reload()
        _M.auto_match_all(cx, _eng)
        F.backfill_owners(cx)
    except Exception:
        pass
    F.rebuild_stock(cx)
    for y in years:
        F.validate_year(cx, y, owner_tin)

    # Hisob-kitob Excel formulalarida (1.6.0) - ilova faqat ma'lumotni yozadi
    return build_formula_workbook(cx, out_path, years, owner_name, owner_tin, progress)


def default_filename(years, owner_name=None):
    ys = sorted(y for y in (years or []) if y)
    if not ys:
        rng = "hisobot"
    elif len(ys) == 1:
        rng = "%d" % ys[0]
    else:
        rng = "%d-%d" % (ys[0], ys[-1])
    tail = ""
    if owner_name:
        # fayl nomiga tashkilotning birinchi ikki so'zi (Windows taqiqlagan belgilarsiz)
        tail = " " + " ".join(re.sub(r'[\\/:*?"<>|]+', " ", owner_name).split()[:2])
    return "КАМЕРАЛ ТЕКШИРУВЛАР %s%s.xlsx" % (rng, tail)


# ===========================================================================
# FORMULALI HISOBOT (1.6.0)
#
# Egasi: "excellarni hisoblashni formula orqali qilsin - ilova sen qilib berma,
# ilova excel formulalaridan foydalanib hisoblasin". Shuning uchun:
#   Ma'lumot     - QQS va ustama SARIQ katakda (nomli: QQS, USTAMA_YYYY); ularni
#                  o'zgartirsa - butun fayl qayta hisoblanadi; oylar jadvali
#                  (COUNTIFS / SUMIFS); jamilar ТХ ga havola.
#   Фактуралар   - har faktura bir satr: satrlar soni va summasi - formula.
#   Кирим        - har faktura satri: summa = miqdor x narx (formula).
#   касса YYYY   - har chek satri + "Товар коди" (ilova bog'lagan tovar).
#   YYYY ТХ      - asl 22 ustun, HAMMA son formula:
#                    kirim  = INDEX(Кирим), narx = INDEX(Кирим)
#                    chiqim = SUMIFS(тақсимот YYYY) shu partiya bo'yicha
#                    qoldiq = boshiga + kirim - chiqim
#   тақсимот YYYY - har sotuv qaysi partiyadan chiqqani (1.7.0, buxgalter qoidasi):
#                  sotuv sanasigacha kelgan ENG OXIRGI fakturadan (_allocate)
#                    boshiga qoldiq = o'tgan yil ТХ dagi shu partiyaning qoldig'i
#                    sotish narxi = tannarx x (1+ustama) x (1+QQS)
#   Xatolar      - oddiy so'z bilan.
# Formulada qilib bo'lmaydigan ish - chekdagi sotuvni qaysi tovarga bog'lash
# (bh_matching) va qaysi partiyadan yechish (_allocate); natijasi "Товар коди"
# va "тақсимот" varag'ida ochiq turibdi.
# Excel faylni ochganda hamma formulani o'zi hisoblaydi (fullCalcOnLoad).
# ===========================================================================
MONTHS_UZ = ["Январь", "Февраль", "Март", "Апрель", "Май", "Июнь", "Июль", "Август",
             "Сентябрь", "Октябрь", "Ноябрь", "Декабрь"]
SH_INFO = "Ma'lumot"
SH_FAKT = "Фактуралар"
SH_KIRIM = "Кирим"
NAME_VAT = "QQS"


def _txt(v):
    """Ma'lumot matni "=" bilan boshlansa Excel uni formula deb o'qimasin."""
    return (" " + v) if isinstance(v, str) and v.startswith("=") else v


def _r4(x):
    return round(float(x or 0) + 0.0, 4)


def _q(sheet):
    return "'%s'" % sheet.replace("'", "''")


def _kassa_name(y):
    return "касса %d" % y


def _alloc_name(y):
    return "тақсимот %d" % y


def _tx_name(y):
    return "%d ТХ" % y


def _mk_name(y):
    return "USTAMA_%d" % y


def _gkey(g):
    return "Т%05d" % g


def _ukey(i):
    return "Б%04d" % i


def _formula_data(cx, owner_tin, years):
    """Bazadan formulali hisobot uchun xom ma'lumot (tanlangan tashkilot)."""
    import bh_matching as M

    roots = M.load_product_groups(cx)

    def root(pid):
        return None if pid is None else roots.get(pid, pid)

    own = (owner_tin, owner_tin)
    lines = [dict(r) for r in cx.execute("""
        SELECT l.id, l.kind, l.document_id doc_id, l.line_no, l.raw_name, l.norm_name, l.mxik,
               l.mxik_name, l.barcode, l.marking_code, l.is_marked, l.unit_raw,
               CAST(l.qty AS REAL) qty, CAST(l.unit_price_net AS REAL) price,
               CAST(l.vat_amount AS REAL) vat, CAST(l.amount_gross AS REAL) gross,
               l.product_id, l.match_method, COALESCE(l.line_date, d.doc_date) dt,
               d.doc_no, d.doc_date, d.partner_name, d.partner_tin, d.pos_id, d.check_type,
               d.is_return, CAST(d.total_gross AS REAL) doc_gross, CAST(d.total_vat AS REAL) doc_vat,
               sf.filename
        FROM doc_line l JOIN document d ON d.id = l.document_id
        LEFT JOIN source_file sf ON sf.id = d.source_file_id
        WHERE (? IS NULL OR d.owner_tin = ?) AND COALESCE(l.line_date, d.doc_date) IS NOT NULL
        ORDER BY COALESCE(l.line_date, d.doc_date), d.id, l.line_no, l.id""", own)]
    last = max(years)
    lines = [l for l in lines if int(l["dt"][:4]) <= last]
    for l in lines:
        l["group"] = root(l["product_id"])
    kirim = [l for l in lines if l["kind"] == "kirim"]
    sales = [l for l in lines if l["kind"] == "chiqim"]

    # bog'lanmagan sotuvlar: nom bo'yicha kalit, kattasidan
    ug = {}
    for s in sales:
        if s["group"] is None:
            g = ug.setdefault(s["norm_name"], {"gross": 0.0})
            g["gross"] += s["gross"] or 0
    ukeys = {n: _ukey(i + 1) for i, (n, _g) in
             enumerate(sorted(ug.items(), key=lambda kv: -kv[1]["gross"]))}
    for s in sales:
        s["key"] = _gkey(s["group"]) if s["group"] is not None else ukeys.get(s["norm_name"], "")
    for k in kirim:
        k["key"] = _gkey(k["group"]) if k["group"] is not None else ""

    # fakturalar (hujjat bo'yicha)
    fakt, seen = [], set()
    for k in kirim:
        if k["doc_id"] in seen or not k["doc_date"]:
            continue
        seen.add(k["doc_id"])
        fakt.append(k)
    return {"kirim": kirim, "sales": sales, "fakt": fakt}


A_BEFORE = "Sotuvgacha kelgan faktura"
A_AFTER = "Sotuvdan keyin kelgan faktura"
A_SHORT = "Omborda yetmadi"
A_NOLOT = "Fakturasi topilmadi"
A_RETURN = "Qaytarildi"


def _allocate(kirim, sales):
    """
    Har sotuv qaysi partiya (faktura satri)dan chiqqani - buxgalter qoidasi (1.7.0).
    Ilgari yillik FIFO edi: 28.09 da sotilgan 45 ta mototsikl avval avgustdagi
    43 talik partiyadan yechilib, 24.09 dagi 45 talik fakturaga 2 ta qolardi.
    Endi:
      a) sotuv sanasigacha kelgan partiyalardan - ENG OXIRGISIDAN boshlab,
         u tugasa oldingisidan;
      b) yetmasa - shu yil ichida sotuvdan keyin kelgan partiyadan (yaqinidan);
      c) baribir yetmasa - partiyasiz ("omborda yetmadi").
    Qaytarish (manfiy miqdor) oxirgi yechilgan partiyaga qaytadi.
    Natija: {sotuv satri id: [(partiya yoki None, miqdor, izoh)]}.
    """
    lots = {}
    for k in sorted(kirim, key=lambda k: (k["dt"], k["id"])):
        if k["group"] is not None and (k["qty"] or 0) > 0:
            lots.setdefault(k["group"], []).append(k)
    left = {k["id"]: _r4(k["qty"]) for gl in lots.values() for k in gl}
    taken = {}      # guruh -> [[partiya, miqdor]] yechilgan tartibda (qaytarish uchun)
    out = {}
    for s in sales:                     # _formula_data sana tartibida beradi
        g, q = s["group"], _r4(s["qty"])
        if g is None or not q:
            continue
        gl = lots.get(g, [])
        res = out.setdefault(s["id"], [])
        if q < 0:
            back, st = -q, taken.get(g, [])
            while back > 0 and st:
                lot, n = st[-1]
                b = min(n, back)
                left[lot["id"]] = _r4(left[lot["id"]] + b)
                res.append((lot, -b, A_RETURN))
                back = _r4(back - b)
                if b == n:
                    st.pop()
                else:
                    st[-1][1] = _r4(n - b)
            if back > 0:
                res.append((None, -back, A_RETURN))
            continue
        need, y = q, s["dt"][:4]
        for cand, why in (([l for l in reversed(gl) if l["dt"] <= s["dt"]], A_BEFORE),
                          ([l for l in gl if l["dt"] > s["dt"] and l["dt"][:4] == y], A_AFTER)):
            for l in cand:
                if need <= 0:
                    break
                av = left[l["id"]]
                if av <= 0:
                    continue
                t = min(av, need)
                left[l["id"]] = _r4(av - t)
                need = _r4(need - t)
                res.append((l, t, why))
                taken.setdefault(g, []).append([l, t])
        if need > 0:
            res.append((None, need, A_SHORT if gl else A_NOLOT))
    return out


def _tx_rows(kirim, sales, year, markup, alloc):
    """
    Bir yil uchun ТХ satrlari - Excel formulasi bilan AYNAN bir xil hisob (satrlar
    ro'yxatini tuzish va jamilarni ilova oynasida ko'rsatish uchun).
    alloc - _allocate natijasi.
    """
    lots = sorted([k for k in kirim if k["group"] is not None and (k["qty"] or 0) > 0
                   and int(k["dt"][:4]) <= year], key=lambda k: (k["dt"], k["id"]))
    out_y, short = {}, {}           # (partiya id, yil) / (guruh, yil) -> miqdor
    for s in sales:
        y = int(s["dt"][:4])
        for lot, q, _why in alloc.get(s["id"], ()):
            k, d = ((lot["id"], y), out_y) if lot is not None else ((s["group"], y), short)
            d[k] = _r4(d.get(k, 0) + q)
    rows, last_lot = [], {}
    for l in lots:
        ly = int(l["dt"][:4])
        op = _r4(l["qty"] - sum(out_y.get((l["id"], y), 0) for y in range(ly, year))) if ly < year else 0
        inq = l["qty"] if ly == year else 0
        out = out_y.get((l["id"], year), 0)
        last_lot[l["group"]] = l
        if op or inq or out:
            rows.append({"kind": "lot", "lot": l, "key": _gkey(l["group"]), "date": l["dt"],
                         "name": l["raw_name"], "cost": l["price"] or 0,
                         "open": _r4(op), "in": _r4(inq), "out": _r4(out)})
    for (g, y), need in sorted(short.items()):
        if y == year and g in last_lot and need:
            lt = last_lot[g]
            rows.append({"kind": "short", "lot": lt, "key": _gkey(g), "date": "%d-12-31" % year,
                         "name": lt["raw_name"], "cost": lt["price"] or 0,
                         "open": 0, "in": 0, "out": _r4(need)})
    with_lots = set(last_lot)
    other = {}
    for s in sales:
        if int(s["dt"][:4]) != year:
            continue
        if s["group"] is None or s["group"] not in with_lots:
            other.setdefault(s["key"], []).append(s)
    for key, ss in other.items():
        q = _r4(sum(s["qty"] or 0 for s in ss))
        if not q:
            continue
        g = _r4(sum(s["gross"] or 0 for s in ss))
        cost = _r4(g / q / (1 + float(C.VAT_RATE)) / (1 + markup))
        f = min(ss, key=lambda s: s["dt"])
        rows.append({"kind": "unmatched", "lot": f, "key": key, "date": f["dt"],
                     "name": f["raw_name"], "cost": cost, "open": 0, "in": 0, "out": q})
    rank = {"lot": 0, "short": 1, "unmatched": 2}
    rows.sort(key=lambda r: (r["date"], rank[r["kind"]], r["name"] or ""))
    for r in rows:
        r["close"] = _r4(r["open"] + r["in"] - r["out"])
        r["sale_price"] = _r4(r["cost"] * (1 + markup) * (1 + float(C.VAT_RATE)))
    return rows


def _style_cells(ws, row, cols, numfmt=None, font=None, fill=None, border=None, align=None):
    for c in cols:
        cell = ws.cell(row=row, column=c)
        if numfmt:
            cell.number_format = numfmt
        if font is not None:
            cell.font = font
        if fill is not None:
            cell.fill = fill
        if border is not None:
            cell.border = border
        if align is not None:
            cell.alignment = align


def build_formula_workbook(cx, out_path, years, owner_name, owner_tin=None, progress=None):
    """Formulali «КАМЕРАЛ ТЕКШИРУВЛАР». (yo'l, statistika) qaytaradi."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.workbook.defined_name import DefinedName

    years = sorted(years)
    year = years[-1]
    data = _formula_data(cx, owner_tin, years)
    kirim, sales = data["kirim"], data["sales"]
    alloc = _allocate(kirim, sales)
    vat = float(C.VAT_RATE)
    markups = {y: float(DB.get_markup(cx, y)) for y in years}

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    try:
        wb.calculation.fullCalcOnLoad = True
    except Exception:
        pass
    S = StyleBook(wb)

    thin = Side(style="thin", color="9AA0A6")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    f_hdr = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    fill_hdr = PatternFill("solid", fgColor="44546A")
    fill_sub = PatternFill("solid", fgColor="5B7397")
    fill_tot = PatternFill("solid", fgColor="D9E2F3")
    fill_err = PatternFill("solid", fgColor="FCE4E4")
    fill_warn = PatternFill("solid", fgColor="FFF2CC")
    fill_in = PatternFill("solid", fgColor="FFF9DB")
    a_hdr = Alignment(horizontal="center", vertical="center", wrap_text=True)
    f_bold = Font(name="Calibri", size=10, bold=True)

    def hdr(ws, r, c, text, sub=False):
        cell = ws.cell(row=r, column=c, value=text)
        cell.font, cell.alignment, cell.border = f_hdr, a_hdr, box
        cell.fill = fill_sub if sub else fill_hdr

    info = wb.create_sheet(SH_INFO)
    fakt = wb.create_sheet(SH_FAKT)
    kir = wb.create_sheet(SH_KIRIM)
    for y in years:
        wb.create_sheet(_kassa_name(y))
    for y in years:
        wb.create_sheet(_tx_name(y))
    for y in years:
        wb.create_sheet(_alloc_name(y))

    # ------------------------------------------------ Ma'lumot: sozlama katakchalari
    for col, w in zip("ABCDEFG", (3, 44, 22, 22, 22, 22, 22)):
        info.column_dimensions[col].width = w
    info.merge_cells("B1:G1")
    info["B1"] = "%s - tovar hisoboti (КАМЕРАЛ ТЕКШИРУВЛАР)" % (owner_name or "Ташкилот")
    info["B1"].font = Font(size=16, bold=True, color="1F3864")
    info["B3"], info["C3"] = "STIR", owner_tin or "hammasi"
    info["B4"], info["C4"] = "Tayyorlangan", datetime.datetime.now().strftime("%d.%m.%Y %H:%M")
    info["B5"], info["C5"] = "Dastur versiyasi", C.VERSION
    info["B7"] = "Sozlamalar (sariq katakni o'zgartirsangiz - butun fayl qayta hisoblanadi)"
    info["B7"].font = Font(size=12, bold=True)
    info["B8"], info["C8"] = "QQS stavkasi", vat
    info["C8"].number_format, info["C8"].fill, info["C8"].border = "0.00%", fill_in, box
    wb.defined_names[NAME_VAT] = DefinedName(NAME_VAT, attr_text="%s!$C$8" % _q(SH_INFO))
    r = 9
    for y in years:
        info.cell(row=r, column=2, value="Ustama %d yil (tannarxga qo'shiladi)" % y)
        c = info.cell(row=r, column=3, value=markups[y])
        c.number_format, c.fill, c.border = "0.00%", fill_in, box
        wb.defined_names[_mk_name(y)] = DefinedName(_mk_name(y), attr_text="%s!$C$%d" % (_q(SH_INFO), r))
        r += 1
    info_tot_row = r + 1

    # ------------------------------------------------ Фактуралар
    for col, w in zip("ABCDEFGH", (6, 12, 18, 44, 16, 10, 18, 52)):
        fakt.column_dimensions[col].width = w
    fakt["A1"] = "Qo'shilgan hisob-fakturalar - har biri bir satr"
    fakt["A1"].font = Font(size=14, bold=True)
    for i, h in enumerate(["№", "Sana", "Faktura raqami", "Yetkazib beruvchi", "STIR", "Satrlar",
                           "Summa (QQS bilan)", "Fayl nomi"]):
        hdr(fakt, 3, i + 1, h)
    kq = _q(SH_KIRIM)
    fr = 4
    for i, d in enumerate(data["fakt"]):
        fakt.cell(row=fr, column=1, value=i + 1)
        fakt.cell(row=fr, column=2, value=C.parse_date(d["doc_date"])).number_format = NUM_DATE
        fakt.cell(row=fr, column=3, value=_txt(d["doc_no"]))
        fakt.cell(row=fr, column=4, value=_txt(d["partner_name"]))
        fakt.cell(row=fr, column=5, value=d["partner_tin"])
        fakt.cell(row=fr, column=6, value="=COUNTIFS(%s!$B:$B,C%d,%s!$D:$D,E%d)" % (kq, fr, kq, fr))
        fakt.cell(row=fr, column=7, value="=SUMIFS(%s!$L:$L,%s!$B:$B,C%d,%s!$D:$D,E%d)" % (kq, kq, fr, kq, fr)
                  ).number_format = NUM_MONEY
        fakt.cell(row=fr, column=8, value=_txt(d["filename"]))
        _style_cells(fakt, fr, range(1, 9), border=box)
        fr += 1
    fakt.cell(row=fr, column=4, value="JAMI faktura:")
    fakt.cell(row=fr, column=5, value="=COUNTA(C4:C%d)" % (fr - 1))
    fakt.cell(row=fr, column=6, value="=SUM(F4:F%d)" % (fr - 1))
    fakt.cell(row=fr, column=7, value="=SUM(G4:G%d)" % (fr - 1)).number_format = NUM_MONEY
    _style_cells(fakt, fr, range(1, 9), font=f_bold, fill=fill_tot)
    fakt.freeze_panes = "A4"
    if fr > 4:
        fakt.auto_filter.ref = "A3:H%d" % (fr - 1)

    # ------------------------------------------------ Кирим
    for col, w in zip("ABCDEFGHIJKLMN", (11, 18, 34, 14, 50, 20, 14, 11, 16, 18, 16, 18, 10, 10)):
        kir.column_dimensions[col].width = w
    kir["A1"] = "Kirim - hisob-fakturalardagi har bir tovar satri (narx QQS siz)"
    kir["A1"].font = Font(size=14, bold=True)
    for i, h in enumerate(["Sana", "Faktura raqami", "Yetkazib beruvchi", "STIR", "Tovar nomi", "MXIK",
                           "O'lchov birligi", "Miqdori", "Narxi (QQS siz)", "Summasi (miqdor x narx)",
                           "QQS", "QQS bilan", "Tovar kodi", "Partiya №"]):
        hdr(kir, 3, i + 1, h)
    kr = 4
    for k in kirim:
        kir.cell(row=kr, column=1, value=C.parse_date(k["dt"])).number_format = NUM_DATE
        kir.cell(row=kr, column=2, value=_txt(k["doc_no"]))
        kir.cell(row=kr, column=3, value=_txt(k["partner_name"]))
        kir.cell(row=kr, column=4, value=k["partner_tin"])
        kir.cell(row=kr, column=5, value=_txt(k["raw_name"]))
        kir.cell(row=kr, column=6, value=k["mxik"])
        kir.cell(row=kr, column=7, value=k["unit_raw"])
        kir.cell(row=kr, column=8, value=_r4(k["qty"])).number_format = NUM_QTY
        kir.cell(row=kr, column=9, value=_r4(k["price"])).number_format = NUM_MONEY
        kir.cell(row=kr, column=10, value="=H%d*I%d" % (kr, kr)).number_format = NUM_MONEY
        kir.cell(row=kr, column=11, value=_r4(k["vat"])).number_format = NUM_MONEY
        kir.cell(row=kr, column=12, value="=J%d+K%d" % (kr, kr)).number_format = NUM_MONEY
        kir.cell(row=kr, column=13, value=k["key"])
        kir.cell(row=kr, column=14, value=k["id"])
        kr += 1
    kir.cell(row=kr, column=5, value="JAMI")
    for c, fmt in ((8, NUM_QTY), (10, NUM_MONEY), (11, NUM_MONEY), (12, NUM_MONEY)):
        L = get_column_letter(c)
        kir.cell(row=kr, column=c, value="=SUBTOTAL(9,%s4:%s%d)" % (L, L, kr - 1)).number_format = fmt
    _style_cells(kir, kr, range(1, 15), font=f_bold, fill=fill_tot)
    kir.freeze_panes = "A4"
    if kr > 4:
        kir.auto_filter.ref = "A3:N%d" % (kr - 1)

    # ------------------------------------------------ касса YYYY
    KH = ["Махсулот\nИдси", "СТИР/\nЖИШШР", "ФМ рақами", "Чек санаси", "Чек рақами", "Маҳсулот номи",
          "Миқдори", "Нархи\n(сатр суммаси)", "Чегирма суммаси", "Дисконт суммаси", "ҚҚС суммаси",
          "Жами нақд пул", "Жами банк карта", "Жами ҚҚС", "Маҳсулот коди", "Ўлчов бирлиги коди",
          "Штрих код", "Воситачи СТИРи (ЖИШШРи)", "Маркировка\nкоди", "Чек тури",
          "Tovar kodi\n(ТХ dagi)", "Qanday bog'langan", "Chek\n(1-satr)", "Партия №\n(тақсимот)"]
    KW = [30, 16, 16, 12, 10, 46, 11, 18, 11, 11, 16, 15, 16, 15, 19, 12, 14, 14, 20, 11, 12, 22, 9, 16]
    HOW = {"barcode": "Shtrix-kod bo'yicha", "marking": "Markirovka bo'yicha", "alias": "Avval tanlangan",
           "name": "Nomi bir xil", "prefix": "Nomi (qisqartirilgan)", "fuzzy": "Nomi o'xshash",
           "fuzzy+mxik": "Nomi o'xshash, MXIK bir xil", "mxik": "MXIK bo'yicha", "manual": "Qo'lda tanlangan"}
    kassa_rows = 0
    for y in years:
        ws = wb[_kassa_name(y)]
        for i, w in enumerate(KW):
            ws.column_dimensions[get_column_letter(i + 1)].width = w
        for i, h in enumerate(KH):
            hdr(ws, 1, i + 1, h)
        ws.row_dimensions[1].height = 32
        rr, seen_doc = 2, set()
        for s in sales:
            if int(s["dt"][:4]) != y:
                continue
            first = s["doc_id"] not in seen_doc
            seen_doc.add(s["doc_id"])
            vals = ["", s["partner_tin"], s["pos_id"], C.parse_date(s["dt"]), s["doc_no"], s["raw_name"],
                    _r4(s["qty"]), _r4(s["gross"]), 0, 0, _r4(s["vat"]), 0,
                    _r4(s["doc_gross"]) if first else None, _r4(s["doc_vat"]) if first else None,
                    s["mxik"], s["unit_raw"], s["barcode"], "", s["marking_code"], s["check_type"],
                    s["key"], HOW.get(s["match_method"], "Bog'lanmagan") if s["group"] is not None
                    else "Bog'lanmagan", 1 if first else None,
                    ", ".join(str(a[0]["id"]) for a in alloc.get(s["id"], ()) if a[0] is not None) or None]
            for i, v in enumerate(vals):
                ws.cell(row=rr, column=i + 1, value=_txt(v))
            ws.cell(row=rr, column=4).number_format = NUM_DATE
            ws.cell(row=rr, column=7).number_format = NUM_QTY
            for c in (8, 11, 12, 13, 14):
                ws.cell(row=rr, column=c).number_format = NUM_MONEY
            if s["is_return"]:
                _style_cells(ws, rr, range(1, 25), fill=fill_warn)
            elif s["group"] is None:
                _style_cells(ws, rr, (21, 22), fill=fill_err)
            rr += 1
            kassa_rows += 1
        ws.cell(row=rr, column=6, value="JAMI")
        for c, fmt in ((7, NUM_QTY), (8, NUM_MONEY), (11, NUM_MONEY), (23, "0")):
            L = get_column_letter(c)
            ws.cell(row=rr, column=c, value="=SUBTOTAL(9,%s2:%s%d)" % (L, L, max(2, rr - 1))).number_format = fmt
        _style_cells(ws, rr, range(1, 25), font=f_bold, fill=fill_tot)
        ws.freeze_panes = "A2"
        if rr > 2:
            ws.auto_filter.ref = "A1:X%d" % (rr - 1)

    # ------------------------------------------------ тақсимот YYYY: sotuv qaysi partiyadan
    AH = ["Чек санаси", "Чек рақами", "Маҳсулот номи", "Миқдори\n(шу партиядан)", "Партия №",
          "Партия санаси", "Фактура рақами", "Калит\n(ТХ даги)", "Изоҳ"]
    AW = [12, 10, 46, 14, 11, 12, 16, 12, 30]
    for y in years:
        ws = wb[_alloc_name(y)]
        for i, w in enumerate(AW):
            ws.column_dimensions[get_column_letter(i + 1)].width = w
        for i, h in enumerate(AH):
            hdr(ws, 1, i + 1, h)
        ws.row_dimensions[1].height = 32
        rr = 2
        for s in sales:
            if int(s["dt"][:4]) != y or not s["qty"]:
                continue
            parts = alloc.get(s["id"]) or [(None, _r4(s["qty"]), A_NOLOT)]
            for lot, q, why in parts:
                vals = [C.parse_date(s["dt"]), s["doc_no"], s["raw_name"], q,
                        lot["id"] if lot is not None else None,
                        C.parse_date(lot["dt"]) if lot is not None else None,
                        lot["doc_no"] if lot is not None else None,
                        lot["id"] if lot is not None else s["key"], why]
                for i, v in enumerate(vals):
                    ws.cell(row=rr, column=i + 1, value=_txt(v))
                ws.cell(row=rr, column=1).number_format = NUM_DATE
                ws.cell(row=rr, column=6).number_format = NUM_DATE
                ws.cell(row=rr, column=4).number_format = NUM_QTY
                if lot is None:
                    _style_cells(ws, rr, range(1, 10), fill=fill_err)
                elif why == A_AFTER:
                    _style_cells(ws, rr, range(1, 10), fill=fill_warn)
                rr += 1
        ws.cell(row=rr, column=3, value="JAMI")
        ws.cell(row=rr, column=4, value="=SUBTOTAL(9,D2:D%d)" % max(2, rr - 1)).number_format = NUM_QTY
        _style_cells(ws, rr, range(1, 10), font=f_bold, fill=fill_tot)
        ws.freeze_panes = "A2"
        if rr > 2:
            ws.auto_filter.ref = "A1:I%d" % (rr - 1)

    # ------------------------------------------------ YYYY ТХ
    TH = ["№", "санаси", "Товар номи", "Маркировкаланган", "МХИК Коди", "Ўлчов бирлиги",
          "1 бирликни сотиш нархи\n(таннарх+устама+ҚҚС)"]
    TG = [(8, "давр бошига қолдиқ (таннарх)"), (11, "Кирим (таннарх)"), (14, "Чиқим (таннарх)"),
          (17, "давр охирига қолдиқ таннархда"), (20, "Чиқим сотиш нархида ҚҚС билан")]
    TW = [6, 11, 46, 16, 40, 16, 16, 12, 16, 19, 12, 16, 19, 12, 16, 19, 12, 16, 19, 12, 16, 19, 11, 9]
    stats = {"years": {}, "kassa_rows": kassa_rows, "tx_rows": 0}
    tot_cells = {}
    for yi, y in enumerate(years):
        if progress:
            progress(yi, len(years), "%d ТХ" % y)
        ws = wb[_tx_name(y)]
        ks = _q(_kassa_name(y))
        aq = _q(_alloc_name(y))
        prev = _q(_tx_name(y - 1)) if (y - 1) in years else None
        mk = _mk_name(y)
        rows = _tx_rows(kirim, sales, y, markups[y], alloc)
        for i, w in enumerate(TW):
            ws.column_dimensions[get_column_letter(i + 1)].width = w
        ws.merge_cells(start_row=1, start_column=2, end_row=1, end_column=22)
        ws.cell(row=1, column=2, value="%s ning %d йил товар хисоботи" % (owner_name or "Ташкилот", y)).font = \
            Font(size=14, bold=True, color="1F3864")
        for i, h in enumerate(TH):
            ws.merge_cells(start_row=3, start_column=i + 1, end_row=4, end_column=i + 1)
            hdr(ws, 3, i + 1, h)
        for c, title in TG:
            ws.merge_cells(start_row=3, start_column=c, end_row=3, end_column=c + 2)
            hdr(ws, 3, c, title)
            for j, sub in enumerate(("Микдори", "Нархи", "суммаси")):
                hdr(ws, 4, c + j, sub, sub=True)
        for c, title in ((23, "Tovar kodi"), (24, "Partiya №")):
            ws.merge_cells(start_row=3, start_column=c, end_row=4, end_column=c)
            hdr(ws, 3, c, title)
        ws.row_dimensions[3].height = 34

        first_r = 5
        for i, x in enumerate(rows):
            rr = first_r + i
            lot = x["lot"]
            is_lot = x["kind"] == "lot"
            this_year = is_lot and x["date"][:4] == str(y)
            sold = "SUMIFS(%s!$G:$G,%s!$U:$U,$W%d)" % (ks, ks, rr)
            name = x["name"] or ""
            if x["kind"] == "unmatched":
                name += "  (fakturasi topilmadi)"
            elif x["kind"] == "short":
                name += "  (omborda yetmadi)"
            mark = lot.get("is_marked")
            ws.cell(row=rr, column=1, value=i + 1)
            ws.cell(row=rr, column=2, value=C.parse_date(x["date"])).number_format = NUM_DATE
            ws.cell(row=rr, column=3, value=_txt(name))
            ws.cell(row=rr, column=4, value="маркировкаланган" if mark else ("маркировкасиз" if mark == 0 else ""))
            ws.cell(row=rr, column=5, value=("%s - %s" % (lot.get("mxik") or "", lot.get("mxik_name") or "")).strip(" -"))
            ws.cell(row=rr, column=6, value=lot.get("unit_raw") or "")
            if is_lot:
                ws.cell(row=rr, column=9, value="=INDEX(%s!$I:$I,MATCH($X%d,%s!$N:$N,0))" % (kq, rr, kq))
            elif x["kind"] == "unmatched":
                ws.cell(row=rr, column=9, value="=IFERROR(SUMIFS(%s!$H:$H,%s!$U:$U,$W%d)/%s/(1+%s)/(1+%s),0)"
                        % (ks, ks, rr, sold, NAME_VAT, mk))
            else:
                ws.cell(row=rr, column=9, value=_r4(x["cost"]))
            ws.cell(row=rr, column=7, value="=I%d*(1+%s)*(1+%s)" % (rr, mk, NAME_VAT))
            if is_lot and not this_year:
                ws.cell(row=rr, column=8, value=("=IFERROR(INDEX(%s!$Q:$Q,MATCH($X%d,%s!$X:$X,0)),0)" % (prev, rr, prev))
                        if prev else x["open"])
            else:
                ws.cell(row=rr, column=8, value=0)
            ws.cell(row=rr, column=10, value="=H%d*I%d" % (rr, rr))
            ws.cell(row=rr, column=11, value=("=INDEX(%s!$H:$H,MATCH($X%d,%s!$N:$N,0))" % (kq, rr, kq)) if this_year else 0)
            ws.cell(row=rr, column=12, value="=I%d" % rr)
            ws.cell(row=rr, column=13, value="=K%d*L%d" % (rr, rr))
            # CHIQIM - тақсимот varag'idan: partiya satri - Партия № bo'yicha, qolgani - tovar kodi
            ws.cell(row=rr, column=14, value="=SUMIFS(%s!$D:$D,%s!$H:$H,$%s%d)"
                    % (aq, aq, "X" if is_lot else "W", rr))
            ws.cell(row=rr, column=15, value="=I%d" % rr)
            ws.cell(row=rr, column=16, value="=N%d*O%d" % (rr, rr))
            ws.cell(row=rr, column=17, value="=H%d+K%d-N%d" % (rr, rr, rr))
            ws.cell(row=rr, column=18, value="=I%d" % rr)
            ws.cell(row=rr, column=19, value="=Q%d*R%d" % (rr, rr))
            ws.cell(row=rr, column=20, value="=N%d" % rr)
            ws.cell(row=rr, column=21, value="=G%d" % rr)
            ws.cell(row=rr, column=22, value="=T%d*U%d" % (rr, rr))
            ws.cell(row=rr, column=23, value=x["key"])
            ws.cell(row=rr, column=24, value=lot["id"] if is_lot else None)
            for c in (8, 11, 14, 17, 20):
                ws.cell(row=rr, column=c).number_format = NUM_QTY
            for c in (7, 9, 10, 12, 13, 15, 16, 18, 19, 21, 22):
                ws.cell(row=rr, column=c).number_format = NUM_MONEY
            if not is_lot:
                _style_cells(ws, rr, range(1, 25), fill=fill_err)
            elif x["close"] < 0:
                _style_cells(ws, rr, range(1, 25), fill=fill_warn)

        tr = first_r + len(rows)
        last_r = tr - 1
        ws.cell(row=tr, column=1, value="ЖАМИ")
        for c in (8, 10, 11, 13, 14, 16, 17, 19, 20, 22):
            L = get_column_letter(c)
            ws.cell(row=tr, column=c, value=("=SUBTOTAL(9,%s%d:%s%d)" % (L, first_r, L, last_r)) if rows else 0
                    ).number_format = NUM_QTY if c in (8, 11, 14, 17, 20) else NUM_MONEY
        _style_cells(ws, tr, range(1, 25), font=f_bold, fill=fill_tot)
        for c, label, formula in ((11, "Кирим жами:", "=M%d" % tr), (14, "Чиқим жами:", "=P%d" % tr),
                                  (17, "Қолдиқ:", "=S%d" % tr),
                                  (20, "Соф фойда:", "=V%d/(1+%s)-P%d" % (tr, NAME_VAT, tr))):
            ws.cell(row=2, column=c, value=label).font = f_bold
            cell = ws.cell(row=2, column=c + 1, value=formula)
            cell.number_format, cell.font = NUM_MONEY, f_bold
        ws.freeze_panes = ws.cell(row=5, column=4)
        if rows:
            ws.auto_filter.ref = "A4:X%d" % last_r
        tot_cells[y] = tr
        T = {"in_sum": sum(r["in"] * r["cost"] for r in rows),
             "out_sum": sum(r["out"] * r["cost"] for r in rows),
             "close_sum": sum(r["close"] * r["cost"] for r in rows),
             "sale_sum": sum(r["out"] * r["sale_price"] for r in rows)}
        T = {k: C.money(Decimal(str(round(v, 4)))) for k, v in T.items()}
        T["profit"] = C.money(C.net_from_gross(T["sale_sum"]) - T["out_sum"])
        stats["years"][y] = T
        stats["tx_rows"] += len(rows)

    # ------------------------------------------------ Ma'lumot: jamilar va oylar
    r = info_tot_row
    info.cell(row=r, column=2, value="Yillar bo'yicha (ТХ varag'idan)").font = Font(size=12, bold=True)
    r += 1
    for i, h in enumerate(["Yil", "Kirim (tannarx)", "Chiqim (tannarx)", "Qoldiq (tannarx)",
                           "Sotuv (ТХ, QQS bilan)", "Kassa tushumi (cheklar)"]):
        hdr(info, r, i + 2, h)
    r += 1
    for y in years:
        tx, tr = _q(_tx_name(y)), tot_cells[y]
        info.cell(row=r, column=2, value=y)
        for c, col in ((3, "M"), (4, "P"), (5, "S"), (6, "V")):
            info.cell(row=r, column=c, value="=%s!%s%d" % (tx, col, tr)).number_format = NUM_MONEY
        ks = _q(_kassa_name(y))
        info.cell(row=r, column=7, value='=SUMIFS(%s!$H:$H,%s!$F:$F,"<>JAMI")' % (ks, ks)).number_format = NUM_MONEY
        _style_cells(info, r, range(2, 8), border=box)
        r += 1
    r += 1
    info.cell(row=r, column=2, value="%d yil - oylar bo'yicha (har oy to'liq qo'shilganini shu yerdan "
                                      "tekshiring)" % year).font = Font(size=12, bold=True)
    r += 1
    for i, h in enumerate(["Oy", "Fakturalar soni", "Kirim (QQS bilan)", "Cheklar soni", "Kassa tushumi", "Holati"]):
        hdr(info, r, i + 2, h)
    r += 1
    fq, kyq = _q(SH_FAKT), _q(_kassa_name(year))
    for m in range(1, 13):
        a, b = "DATE(%d,%d,1)" % (year, m), "EOMONTH(DATE(%d,%d,1),0)" % (year, m)
        info.cell(row=r, column=2, value=MONTHS_UZ[m - 1])
        info.cell(row=r, column=3, value='=COUNTIFS(%s!$B:$B,">="&%s,%s!$B:$B,"<="&%s)' % (fq, a, fq, b))
        info.cell(row=r, column=4, value='=SUMIFS(%s!$L:$L,%s!$A:$A,">="&%s,%s!$A:$A,"<="&%s)'
                  % (kq, kq, a, kq, b)).number_format = NUM_MONEY
        info.cell(row=r, column=5, value='=SUMIFS(%s!$W:$W,%s!$D:$D,">="&%s,%s!$D:$D,"<="&%s)' % (kyq, kyq, a, kyq, b))
        info.cell(row=r, column=6, value='=SUMIFS(%s!$H:$H,%s!$D:$D,">="&%s,%s!$D:$D,"<="&%s)'
                  % (kyq, kyq, a, kyq, b)).number_format = NUM_MONEY
        info.cell(row=r, column=7, value='=IF(AND(C%d=0,E%d=0),"ma\'lumot yo\'q",IF(E%d=0,"chek yo\'q",'
                                         'IF(C%d=0,"faktura yo\'q","bor")))' % (r, r, r, r))
        _style_cells(info, r, range(2, 8), border=box)
        r += 1
    r += 1
    for t in ["Qanday o'qish kerak:",
              "  - Bu fayldagi hamma hisob - Excel formulasi. Katakni bosing - formulasi ko'rinadi.",
              "  - Sariq katakdagi QQS yoki ustamani o'zgartirsangiz, ТХ va jamilar o'zi qayta hisoblanadi.",
              "  - Sotuv sotuv sanasigacha kelgan ENG OXIRGI fakturadan yechiladi, u tugasa oldingisidan. "
              "Qaysi sotuv qaysi fakturadan - \"тақсимот\" varag'ida; ТХ \"Чиқим\" o'shandan yig'iladi.",
              "  - Davr boshiga qoldiq o'tgan yil ТХ varag'idagi shu partiyaning qoldig'idan olinadi.",
              "  - Qizil satr: tovar sotilgan, lekin fakturasi topilmagan yoki omborda yetmagan - tannarx taxminiy.",
              "  - \"Tovar kodi\" - dastur chekdagi sotuvni qaysi fakturadagi tovarga bog'lagani."]:
        info.merge_cells(start_row=r, start_column=2, end_row=r, end_column=7)
        info.cell(row=r, column=2, value=t).font = f_bold if t.endswith(":") else Font(size=10)
        r += 1

    # ------------------------------------------------ Xatolar (oddiy so'z bilan)
    stats["issues"] = write_issues_sheet(wb, cx, years, S, owner_tin)

    d = os.path.dirname(os.path.abspath(out_path))
    if d and not os.path.isdir(d):
        os.makedirs(d, exist_ok=True)
    wb.save(out_path)
    wb.close()
    return out_path, stats
